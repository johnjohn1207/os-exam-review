# -*- coding: utf-8 -*-
"""
create_full_word_doc.py
生成「國立中央大學碩士班入學考試（MA02 計算機概論）作業系統歷年考題篩選暨期中考範圍超詳解手冊」
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_document():
    doc = Document()
    
    # 邊界設定 (標準 1 英吋)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    def set_run_font(run, font_name="微軟正黑體", size_pt=10.5, bold=False, italic=False, color_rgb=(45, 55, 72)):
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        for child in list(tcPr):
            if child.tag.endswith('shd'):
                tcPr.remove(child)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_title(text, subtitle_text=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        set_run_font(run, font_name="微軟正黑體", size_pt=22, bold=True, color_rgb=(26, 54, 93))
        
        if subtitle_text:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(20)
            run_sub = p_sub.add_run(subtitle_text)
            set_run_font(run_sub, font_name="微軟正黑體", size_pt=13, bold=False, color_rgb=(74, 85, 104))

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        set_run_font(run, font_name="微軟正黑體", size_pt=16, bold=True, color_rgb=(26, 54, 93))
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, font_name="微軟正黑體", size_pt=13, bold=True, color_rgb=(43, 108, 176))
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_run_font(run, font_name="微軟正黑體", size_pt=11.5, bold=True, color_rgb=(44, 82, 130))
        return p

    def add_body(text="", bold_prefix="", space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        if bold_prefix:
            run_bp = p.add_run(bold_prefix)
            set_run_font(run_bp, font_name="微軟正黑體", size_pt=10, bold=True, color_rgb=(45, 55, 72))
        if text:
            run_t = p.add_run(text)
            set_run_font(run_t, font_name="微軟正黑體", size_pt=10, bold=False, color_rgb=(45, 55, 72))
        return p

    def add_bullet(text="", bold_prefix="", level=0, space_after=3):
        p = doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            run_bp = p.add_run(bold_prefix)
            set_run_font(run_bp, font_name="微軟正黑體", size_pt=10, bold=True, color_rgb=(45, 55, 72))
        if text:
            run_t = p.add_run(text)
            set_run_font(run_t, font_name="微軟正黑體", size_pt=10, bold=False, color_rgb=(45, 55, 72))
        return p

    def add_callout(lines, title=None, border_color="2B6CB0", bg_color="F7FAFC"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        if title:
            run_title = p.add_run(f"【{title}】\n")
            set_run_font(run_title, font_name="微軟正黑體", size_pt=10.5, bold=True, color_rgb=(43, 108, 176))
        for i, line in enumerate(lines):
            if i > 0 or title:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.2
            run = p.add_run(line)
            set_run_font(run, font_name="微軟正黑體", size_pt=9.5, bold=False, color_rgb=(45, 55, 72))
        
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_code(code_lines):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "EDF2F7")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:color="CBD5E0"/><w:left w:val="single" w:sz="18" w:color="718096"/><w:bottom w:val="single" w:sz="6" w:color="CBD5E0"/><w:right w:val="single" w:sz="6" w:color="CBD5E0"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(borders)
        
        for i, line in enumerate(code_lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(45, 55, 72)
            
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_table_custom(col_widths, headers, data, alt_bg="F7FAFC"):
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            hdr_cells[i].width = Inches(col_widths[i])
            set_cell_background(hdr_cells[i], "2B6CB0")
            set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                set_run_font(r, font_name="微軟正黑體", size_pt=9.5, bold=True, color_rgb=(255, 255, 255))
                
        for r_idx, row_data in enumerate(data):
            row = table.add_row()
            r_cells = row.cells
            bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                r_cells[c_idx].text = str(val)
                r_cells[c_idx].width = Inches(col_widths[c_idx])
                set_cell_background(r_cells[c_idx], bg)
                set_cell_margins(r_cells[c_idx], top=80, bottom=80, left=100, right=100)
                p = r_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                if c_idx in [0, 1, 3, 4, 5]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_run_font(r, font_name="微軟正黑體", size_pt=9, bold=False, color_rgb=(45, 55, 72))
                    
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:color="CBD5E0"/><w:bottom w:val="single" w:sz="6" w:color="CBD5E0"/><w:insideH w:val="single" w:sz="4" w:color="E2E8F0"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
        tblPr.append(borders)
        
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(6)

    # ==========================
    # 文件封面與前言
    # ==========================
    add_title("國立中央大學碩士班入學考試（MA02 計算機概論）", "作業系統歷年試題篩選分析暨期中考範圍超詳解手冊（105 ~ 115 學年度）")
    
    add_callout([
        "文件版本：2026 旗艦精裝版",
        "試卷來源：國立中央大學資訊管理學系碩士班入學考試（科目：計算機概論，代碼 MA02），涵蓋 105 至 115 學年度共 11 份完整歷年試題。",
        "本次考試範圍基準：大學部作業系統（Operating Systems）期中考試範圍，精準對齊講義 CH1 ~ CH5 與經典權威教材《Operating System Concepts 10th》（恐龍本）核心考點。",
        "核心涵蓋單元：CH1（作業系統概論與硬體架構）、CH2（行程概念與處理序間通訊 IPC）、CH3（CPU 排程演算法與即時排程）、CH4（多執行緒程式設計與多核心平行度）、CH5（死結特徵、預防、避免、銀行家演算法、偵測與復原），以及進階高頻同步考點（Peterson 演算法、硬體原子指令、臨界區間三大要件）。",
        "期末範圍排除：恐龍本 CH8 ~ CH19（記憶體管理、虛擬記憶體、檔案系統、磁碟排程、保護安全）以及 Linux 系統操作指令、計概其他科目（網路、資料庫、軟工、OOP 等），於本手冊中清楚列示分類以利後續備考。"
    ], title="研究手冊基本資訊與範圍界定說明")
    
    # ==========================
    # 第一章：試題篩選與範圍對照總覽
    # ==========================
    add_heading_1("第一章 試題篩選原則與學科領域分類評估")
    
    add_heading_2("1.1 評估準則與分類邏輯")
    add_body("中央大學資訊管理學系碩士班入學考試之「計算機概論（MA02）」考科為整合型試卷，命題範疇橫跨作業系統（Operating Systems）、計算機網路（Computer Networks）、資料庫系統（Database Systems）、軟體工程與專案管理（Software Engineering & Project Management）以及物件導向程式設計（Object-Oriented Programming）。")
    add_body("為精準支援本次作業系統期中考試複習，本手冊採取三層式篩選機制：")
    add_bullet("研判試題核心考點是否屬於作業系統理論範疇。若為網路協定（TCP/IP、VPN）、資料庫正規化（1NF~BCNF）、SQL 語法、軟工生命週期或 Java 語法，判定為「非 OS 範疇」。", bold_prefix="第一層：學門過濾。")
    add_bullet("若屬於作業系統題目，進一步比對是否落於本次期中考範圍（CH1 ~ CH5 及臨界區間考點）。若涉及分頁（Paging）、TLB、虛擬記憶體置換（LRU）、檔案系統架構或 RAID 磁碟陣列，判定為「屬於 OS 但超出本次範圍（期末範圍）」。", bold_prefix="第二層：範圍檢驗。")
    add_bullet("凡通過前兩層檢驗之題目，全量收錄至第二章，提供完整中英文對照試題、破題思路、逐步計算推導、甘特圖模擬與標準解答。", bold_prefix="第三層：深度解題。")

    add_heading_2("1.2 歷年考題分類總表（105 ~ 115 學年度全真試題矩陣）")
    add_body("下表彙整 105 至 115 學年度共 11 份試卷之作業系統與相關試題評估結果：")
    
    headers = ["學年度", "題號", "試題核心主題", "所屬領域", "屬於 OS", "期中範圍", "章節與處置"]
    widths = [0.7, 0.6, 2.0, 0.9, 0.7, 0.7, 1.4]
    
    table_data = [
        ["105", "第 2 題", "快取記憶體概念與寫入策略 (Write policy)", "OS / 計算機架構", "是", "符合", "CH1 快取與儲存階層 (詳解)"],
        ["106", "第 2 題", "單核心架構 vs. 微核心架構 (Kernel)", "作業系統核心", "是", "符合", "CH1/CH2 系統結構 (詳解)"],
        ["107", "第 2 題", "存取控制機制 (ACL vs. Capability)", "作業系統防護", "是", "超出", "CH17 保護安全 (期末範圍)"],
        ["107", "第 3 題", "雲端虛擬化 (Hypervisor vs. Container)", "作業系統虛擬化", "是", "超出", "CH18 虛擬化技術 (期末範圍)"],
        ["108", "第 7 題", "銀行家演算法安全狀態與 Request 判定", "作業系統死結", "是", "符合", "CH5 死結避免 (詳解)"],
        ["108", "第 8 題", "TLB 命中率與單/雙階層分頁表 EAT 計算", "記憶體管理", "是", "超出", "CH9 記憶體分頁 (期末範圍)"],
        ["108", "第 9 題", "邏輯位址空間與實體位址位元數推算", "記憶體管理", "是", "超出", "CH8 主記憶體 (期末範圍)"],
        ["108", "第 10 題", "二維陣列迴圈走訪之 LRU 分頁錯誤推算", "虛擬記憶體", "是", "超出", "CH10 虛擬記憶體 (期末範圍)"],
        ["109", "第五題", "行程五狀態轉換模型 (Process State)", "行程管理", "是", "符合", "CH2 行程概念 (詳解)"],
        ["109", "第六題", "死結發生四大必要條件 (Deadlock)", "作業系統死結", "是", "符合", "CH5 死結特性 (詳解)"],
        ["109", "第七題", "Belady 異常現象與 FIFO 分頁置換", "虛擬記憶體", "是", "超出", "CH10 虛擬記憶體 (期末範圍)"],
        ["109", "第八題", "即時 CPU 排程 (RMS 與 EDF 充足性判定)", "CPU 排程", "是", "符合", "CH3 即時排程 (詳解)"],
        ["109", "第九題", "單階層與三階層分頁表 EAT 數值計算", "記憶體管理", "是", "超出", "CH9 記憶體分頁 (期末範圍)"],
        ["110", "第六題", "Peterson 演算法雙行程臨界區間程式填空", "處理序同步", "是", "符合", "CH2/CH6 考點七同步 (詳解)"],
        ["110", "第七題", "核心記憶體配置策略 (Buddy vs. Slab)", "核心記憶體", "是", "超出", "CH9 核心記憶體 (期末範圍)"],
        ["110", "第八題", "TLB 觸及範圍 (TLB Reach) 項目數推算", "記憶體管理", "是", "超出", "CH9 記憶體架構 (期末範圍)"],
        ["111", "第 2 題", "Linux 檔案屬性解讀與 chmod 755 變更", "Linux / 檔案權限", "是", "超出", "CH13 檔案系統 (系統維運)"],
        ["111", "第 3 題", "多重實例資源死結偵測演算法推導", "作業系統死結", "是", "符合", "CH5 死結偵測 (詳解)"],
        ["112", "第 1 題", "Linux crontab 定期排程備份指令設定", "Linux 系統管理", "廣義", "超出", "系統維運操作 (超出範圍)"],
        ["112", "第 3 題", "64 位元電腦 32-bit 行程之分頁表大小 Byte", "記憶體管理", "是", "超出", "CH8 分頁架構 (期末範圍)"],
        ["113", "第 3 題", "TLB Reach、實體位址空間與三階層 EAT", "記憶體管理", "是", "超出", "CH9 記憶體分頁 (期末範圍)"],
        ["113", "第 4 題", "fork() 與 thread_create() 行程與執行緒計數", "行程與多執行緒", "是", "符合", "CH2/CH4 考點七 (詳解)"],
        ["114", "第 11 題", "位元定址分頁系統頁數與實體位址位元", "記憶體管理", "是", "超出", "CH8 主記憶體 (期末範圍)"],
        ["114", "第 12 題", "多層佇列 vs. 多層反饋佇列核心差異", "CPU 排程", "是", "符合", "CH3 CPU 排程 (詳解)"],
        ["114", "第 13 題", "磁碟排程 C-SCAN 與 C-LOOK 移動距離", "次級儲存", "是", "超出", "CH11 磁碟排程 (期末範圍)"],
        ["114", "第 14 題", "巢狀 fork() 與 printf 輸出次數執行追蹤", "行程管理", "是", "符合", "CH2 行程建立 (詳解)"],
        ["114", "第 15 題", "6 顆硬碟 RAID 最小有效資料量評估", "次級儲存", "是", "超出", "CH11 磁碟陣列 (期末範圍)"],
        ["114", "第 17 題", "作業系統開啟檔案內部表格維護機制", "檔案系統", "是", "超出", "CH13 檔案系統 (期末範圍)"],
        ["114", "第 29 題", "行程控制表 (PCB) 收錄內部欄位資訊", "行程管理", "是", "符合", "CH2 行程結構 (詳解)"],
        ["114", "第 30 題", "虛擬記憶體策略之系統優勢研判", "虛擬記憶體", "是", "超出", "CH10 虛擬記憶體 (期末範圍)"],
        ["114", "第 31 題", "臨界區間問題三大必要滿足準則", "處理序同步", "是", "符合", "CH2/CH6 考點七 (詳解)"],
        ["114", "第 33 題", "死結與活結特性、資源排序防範手段", "作業系統死結", "是", "符合", "CH5 死結與活結 (詳解)"],
        ["115", "第 21 題", "死結預防打破條件之活結風險與避免代價", "作業系統死結", "是", "符合", "CH5 死結策略 (詳解)"],
        ["115", "第 22 題", "硬體原子指令、關閉中斷與自旋鎖", "處理序同步", "是", "符合", "CH2/CH6 考點七 (詳解)"],
        ["115", "第 23 題", "連續配置、分頁與分段之內部碎裂問題", "記憶體管理", "是", "超出", "CH8 記憶體配置 (期末範圍)"],
        ["115", "第 24 題", "RAID 0, 1, 5, 6, 0+1 等級特性辨析", "次級儲存", "是", "超出", "CH11 儲存架構 (期末範圍)"],
        ["115", "第 25 題", "即時 CPU 排程 (RMS 週期優先權與 EDF 100%)", "CPU 排程", "是", "符合", "CH3 即時排程 (詳解)"],
        ["115", "第 27 題", "行程與執行緒同步、記憶體共享與 IPC 開銷", "行程與多執行緒", "是", "符合", "CH2/CH4 核心 (詳解)"],
        ["115", "第 28 題", "2TB 磁碟點陣圖 (Bitmap) 管理 Free Space", "檔案系統", "是", "超出", "CH14 檔案實作 (期末範圍)"],
        ["115", "問答 1 題", "TLB 搭配雜湊分頁表 (Hashed Table) 架構與 EAT", "記憶體管理", "是", "超出", "CH9 分頁架構 (期末範圍)"]
    ]
    add_table_custom(widths, headers, table_data)

    add_heading_2("1.3 屬於作業系統但超出本次期中考範圍之題目清單（期末備考備查）")
    add_body("為利日後準備期末考試，特將本次被排除之 OS 試題依恐龍本經典單元歸納如下：")
    add_bullet("包含 108 年第 8 題、第 9 題；109 年第 9 題；110 年第七題（Buddy/Slab）、第八題（TLB Reach）；112 年第 3 題；113 年第 3 題；114 年第 11 題；115 年第 23 題（內部碎裂）、問答第 1 題（雜湊分頁表架構與 EAT）。", bold_prefix="主記憶體與分頁機制（恐龍本 CH8 / CH9）：")
    add_bullet("包含 108 年第 10 題（2D 陣列 LRU 缺頁計算）；109 年第 7 題（Belady's Anomaly）；114 年第 30 題（虛擬記憶體優勢）。", bold_prefix="虛擬記憶體管理（恐龍本 CH10）：")
    add_bullet("包含 114 年第 13 題（C-SCAN / C-LOOK 磁碟臂移動計算）；114 年第 15 題、115 年第 24 題（RAID 0/1/5/6/0+1 特性與空間利用率）。", bold_prefix="次級儲存與磁碟排程（恐龍本 CH11）：")
    add_bullet("包含 114 年第 17 題（System-wide vs. Per-process 開啟檔案表）；115 年第 28 題（2TB 磁碟點陣圖 Free Space 計算）。", bold_prefix="檔案系統介面與實作（恐龍本 CH13 / CH14）：")
    add_bullet("包含 107 年第 2 題（Access Control: ACL vs. Capability-based，CH17）；107 年第 3 題（虛擬化 Hypervisor vs. Container，CH18）；115 年第 26 題（拜占庭將軍問題與分散式容錯，CH19）。", bold_prefix="保護、安全與虛擬化（恐龍本 CH17 / CH18 / CH19）：")
    add_bullet("包含 111 年第 2 題（chmod 755）；112 年第 1 題（crontab 排程設定）；114 年第 18 題（grep）、第 34 題（chmod 權限）、第 35 題（Shell echo 引號）；115 年第 29 題（systemctl）、第 30 題（Linux 管線指令）。", bold_prefix="Linux 系統管理與指令操作：")

    # ==========================
    # 第二章：符合本次期中範圍試題超詳解
    # ==========================
    add_heading_1("第二章 符合期中範圍考題逐題超深度詳解（共 18 道試題）")

    # --- 考題 1: 105 年第 2 題 ---
    add_heading_2("考題一、【105 學年度】第 2 題：快取記憶體概念與寫入策略（Cache Write Policy）")
    add_callout([
        "2. (1) 何謂 Cache？其功用為何？(5%)",
        "   (2) 請比較下列兩種 cache write policy 的優缺點：write-through and write-back。(5%)"
    ], title="試題原始內容（配分 10%）")
    add_body("CH1 作業系統概論與硬體架構（電腦系統組織、儲存階層 Storage Hierarchy、快取機制 Caching）。", bold_prefix="考點定位：")
    add_body("本題評量計算機架構中記憶體階層（Storage Hierarchy）之核心元件 Cache 的功能本質，以及當 CPU 執行寫入動作時，快取與主記憶體保持資料一致性之兩大寫入策略（Write-through 與 Write-back）的工程權衡。", bold_prefix="觀念剖析：")
    add_heading_3("完整詳解與論證：")
    add_bullet("快取記憶體（Cache）是一種位於 CPU 暫存器（Registers）與主記憶體（Main Memory）之間的高速小型靜態隨機存取記憶體（SRAM）。其存取速度遠高於 DRAM，但單位容量成本較高。", bold_prefix="(1) 何謂 Cache 與其功用：")
    add_bullet("Cache 的核心功用在於利用「程式局部性原理（Principle of Locality）」，包含時間局部性（Temporal Locality：近期被存取的資料極可能很快再次被存取）與空間局部性（Spatial Locality：被存取資料相鄰位址的資料極可能很快被存取）。透過將 CPU 頻繁讀寫的指令與資料暫存於快取中，大幅消弭高時脈處理器與慢速主記憶體之間的「速度落差（Speed Gap）」，降低平均記憶體存取時間（Average Memory Access Time, AMAT），提升全系統整體運算吞吐量。", level=1)
    add_bullet("當 CPU 寫入快取且命中時，同時將資料寫入 Cache 與下一層主記憶體（Main Memory）。", bold_prefix="(2) Write-through（直寫式）運作與優缺點：")
    add_bullet("優點：資料一致性極高，主記憶體隨時保有最新資料；控制邏輯簡單，若遇非預期斷電或系統崩潰，復原風險極低。", level=1)
    add_bullet("缺點：寫入延遲顯著，每次寫入速度受限於慢速主記憶體；高頻率佔用系統匯流排頻寬，大幅降低處理器效能（實作上常需額外配置 Write Buffer 緩衝器以緩解阻塞）。", level=1)
    add_bullet("當 CPU 寫入快取且命中時，僅更新 Cache 內部區塊，並將該區塊標記為「髒（Dirty bit = 1）」。唯有當該 Dirty 區塊被置換（Evicted / Replaced）出快取時，才一次性寫回主記憶體。", bold_prefix="Write-back（寫回式 / 回寫式）運作與優缺點：")
    add_bullet("優點：寫入速度維持在快取極速等級；大幅降低對主記憶體匯流排的頻繁存取，顯著節省匯流排頻寬與系統能耗。", level=1)
    add_bullet("缺點：控制邏輯複雜；快取與主記憶體在置換前處於不一致狀態；若在寫回前發生系統當機或斷電，尚未同步的資料將永久遺失。", level=1)

    # --- 考題 2: 106 年第 2 題 ---
    add_heading_2("考題二、【106 學年度】第 2 題：單核心架構 vs. 微核心架構（Kernel Architectures）")
    add_callout([
        "2. Monolithic kernel and microkernel are two main kernel architectures adopted by modern operating systems. Please describe their features and then explain their advantages and disadvantages respectively. (10%)"
    ], title="試題原始內容（配分 10%）")
    add_body("CH1 / CH2 作業系統架構與結構（Operating-System Structures / Monolithic vs. Microkernel）。", bold_prefix="考點定位：")
    add_body("評量現代作業系統核心之兩大主流架構設計哲學：全包式特權架構（Monolithic）與模組化最小特權架構（Microkernel）在特權級別、位址空間、IPC 與容錯性上的根本差異。", bold_prefix="觀念剖析：")
    add_heading_3("完整詳解與論證：")
    add_bullet("將作業系統所有核心服務（行程排程、記憶體管理、檔案系統、網路協定堆疊與硬體驅動程式）全部封裝於單一龐大二進位核心中，且全數執行於最高特權的核心模式（Kernel Mode / Ring 0）與同一定址空間（Single Address Space）內。元件間溝通直接採用高效函式呼叫（Function Calls）。代表系統：Linux、傳統 Unix。", bold_prefix="1. Monolithic Kernel（單核心架構 / 巨集核心）之特徵：")
    add_bullet("優點：執行效能極高。由於所有元件處於同一位址空間，模組間通訊無須透過跨定址空間 IPC，亦無使用者模式與核心模式之間的切換（Mode Switch）開銷。", level=1)
    add_bullet("缺點：可靠性與容錯性低。所有模組皆以最高特權執行，任何單一模組（例如第三方驅動程式）的記憶體越界或指標錯誤，將直接導致全系統崩潰（Kernel Panic / 當機）；此外程式碼高度耦合，維護、擴充與除錯代價極大。", level=1)
    add_bullet("採「極簡化」核心原則，核心模式（Kernel Mode）僅保留最基礎的底層機制（如極簡排程、低階記憶體映射與行程間訊息傳遞 IPC）；將其餘傳統系統服務（如檔案系統、虛擬記憶體、網路協定、裝置驅動程式）移出核心，改以使用者行程（Server Processes）形式運行於受保護的使用者模式（User Mode）。客戶端與伺服元件透過微核心提供的訊息傳遞（Message Passing）進行溝通。代表系統：Mach、QNX、L4。", bold_prefix="2. Microkernel（微核心架構）之特徵：")
    add_bullet("優點：極高的可靠性與安全性。各系統服務在獨立隔離的記憶體空間運行，單一服務崩潰不會波及核心或其他服務，且支援動態重啟；核心極小，易於移植至新硬體架構，且便於模組化擴充。", level=1)
    add_bullet("缺點：效能開銷較大。應用程式請求系統服務時，必須經過微核心進行訊息轉發，伴隨頻繁的使用者模式與核心模式切換，以及多次上下文切換（Context Switch），效能損耗顯著。", level=1)

    # --- 考題 3: 108 年第 7 題 ---
    add_heading_2("考題三、【108 學年度】第 7 題：銀行家演算法（Banker's Algorithm）綜合計算推導")
    add_callout([
        "7. (8 分) Consider the following snapshot of a system:",
        "   Allocation (A B C D), Max (A B C D), Available (A B C D)",
        "   P0: Alloc = (2, 0, 0, 1), Max = (4, 2, 1, 2)",
        "   P1: Alloc = (3, 1, 2, 1), Max = (5, 2, 5, 2)",
        "   P2: Alloc = (2, 1, 0, 3), Max = (2, 3, 1, 6)",
        "   P3: Alloc = (1, 3, 1, 2), Max = (1, 4, 2, 4)",
        "   P4: Alloc = (1, 4, 3, 2), Max = (3, 6, 6, 5)",
        "   Available = (3, 3, 2, 1)",
        "   Answer the following questions using the banker's algorithm:",
        "   (a) If a request from process P1 arrives for (1, 1, 0, 0), can the request be granted immediately? Why?",
        "   (b) If a request from process P4 arrives for (0, 0, 2, 0), can the request be granted immediately? Why?"
    ], title="試題原始內容（配分 8%）")
    add_body("CH5 死結避免（Deadlock Avoidance / Banker's Algorithm 安全性演算法與資源請求演算法）。", bold_prefix="考點定位：")
    add_heading_3("逐步計算推導歷程：")
    add_body("第一步：計算全系統行程之需求矩陣 Need，公式為 Need = Max - Allocation：")
    add_bullet("Need_0 = (4, 2, 1, 2) - (2, 0, 0, 1) = (2, 2, 1, 1)")
    add_bullet("Need_1 = (5, 2, 5, 2) - (3, 1, 2, 1) = (2, 1, 3, 1)")
    add_bullet("Need_2 = (2, 3, 1, 6) - (2, 1, 0, 3) = (0, 2, 1, 3)")
    add_bullet("Need_3 = (1, 4, 2, 4) - (1, 3, 1, 2) = (0, 1, 1, 2)")
    add_bullet("Need_4 = (3, 6, 6, 5) - (1, 4, 3, 2) = (2, 2, 3, 3)")
    add_body("目前初始 Available = (3, 3, 2, 1)。")
    
    add_heading_3("小題 (a) 判定程序：Request_1 = (1, 1, 0, 0)")
    add_bullet("條件檢驗一：Request_1 (1, 1, 0, 0) <= Need_1 (2, 1, 3, 1)，合法。")
    add_bullet("條件檢驗二：Request_1 (1, 1, 0, 0) <= Available (3, 3, 2, 1)，可用資源充足。")
    add_bullet("試探性分配（Pretend Allocation）：")
    add_bullet("Available' = (3, 3, 2, 1) - (1, 1, 0, 0) = (2, 2, 2, 1)", level=1)
    add_bullet("Allocation_1' = (3, 1, 2, 1) + (1, 1, 0, 0) = (4, 2, 2, 1)", level=1)
    add_bullet("Need_1' = (2, 1, 3, 1) - (1, 1, 0, 0) = (1, 0, 3, 1)", level=1)
    add_bullet("執行安全性演算法（Safety Algorithm），令 Work = (2, 2, 2, 1)，Finish 全為 False：")
    add_bullet("檢驗 P0：Need_0 (2, 2, 1, 1) <= Work (2, 2, 2, 1) 成立。P0 執行並釋放資源：Work = (2, 2, 2, 1) + (2, 0, 0, 1) = (4, 2, 2, 2)，Finish[0] = True。", level=1)
    add_bullet("檢驗 P2：Need_2 (0, 2, 1, 3) <= Work (4, 2, 2, 2) 不成立（資源 D: 3 > 2）。跳過。", level=1)
    add_bullet("檢驗 P3：Need_3 (0, 1, 1, 2) <= Work (4, 2, 2, 2) 成立。P3 執行並釋放資源：Work = (4, 2, 2, 2) + (1, 3, 1, 2) = (5, 5, 3, 4)，Finish[3] = True。", level=1)
    add_bullet("檢驗 P1：Need_1' (1, 0, 3, 1) <= Work (5, 5, 3, 4) 成立。P1 執行並釋放資源：Work = (5, 5, 3, 4) + (4, 2, 2, 1) = (9, 7, 5, 5)，Finish[1] = True。", level=1)
    add_bullet("檢驗 P2：Need_2 (0, 2, 1, 3) <= Work (9, 7, 5, 5) 成立。P2 執行並釋放資源：Work = (9, 7, 5, 5) + (2, 1, 0, 3) = (11, 8, 5, 8)，Finish[2] = True。", level=1)
    add_bullet("檢驗 P4：Need_4 (2, 2, 3, 3) <= Work (11, 8, 5, 8) 成立。P4 執行並釋放資源，Finish[4] = True。", level=1)
    add_body("結論 (a)：全體行程 Finish 皆為 True，存在安全序列 <P0, P3, P1, P2, P4>，系統處於安全狀態（Safe State）。因此 P1 的請求可以立即被核准（Yes, granted immediately）。", bold_prefix="判定結論：")

    add_heading_3("小題 (b) 判定程序：Request_4 = (0, 0, 2, 0)")
    add_bullet("條件檢驗一：Request_4 (0, 0, 2, 0) <= Need_4 (2, 2, 3, 3)，合法。")
    add_bullet("條件檢驗二：Request_4 (0, 0, 2, 0) <= Available (3, 3, 2, 1)，可用資源充足。")
    add_bullet("試探性分配：")
    add_bullet("Available'' = (3, 3, 2, 1) - (0, 0, 2, 0) = (3, 3, 0, 1)", level=1)
    add_bullet("Allocation_4'' = (1, 4, 3, 2) + (0, 0, 2, 0) = (1, 4, 5, 2)", level=1)
    add_bullet("Need_4'' = (2, 2, 3, 3) - (0, 0, 2, 0) = (2, 2, 1, 3)", level=1)
    add_bullet("執行安全性演算法，令 Work = (3, 3, 0, 1)：")
    add_bullet("檢驗各行程需求：", level=1)
    add_bullet("Need_0 = (2, 2, 1, 1)：需要 C 資源 1 個，但 Work_C = 0，無法滿足！", level=2)
    add_bullet("Need_1 = (2, 1, 3, 1)：需要 C 資源 3 個，無法滿足！", level=2)
    add_bullet("Need_2 = (0, 2, 1, 3)：需要 C 資源 1 個且 D 需 3 個，無法滿足！", level=2)
    add_bullet("Need_3 = (0, 1, 1, 2)：需要 C 資源 1 個且 D 需 2 個，無法滿足！", level=2)
    add_bullet("Need_4'' = (2, 2, 1, 3)：需要 C 資源 1 個且 D 需 3 個，無法滿足！", level=2)
    add_body("結論 (b)：在試探分配後，系統中沒有任何一個行程的 Need 能夠被 Work 滿足，全體行程皆因資源不足而卡死，無法找到任何安全序列。系統將進入不安全狀態（Unsafe State）。因此 P4 的請求不可以立即被核准（No, cannot be granted immediately），P4 必須等待。", bold_prefix="判定結論：")

    # --- 考題 4: 109 年第五題 ---
    add_heading_2("考題四、【109 學年度】第五題：行程五狀態轉換模型（Process State Diagram）")
    add_callout([
        "第五題: (5 分) Please draw the diagram of process state."
    ], title="試題原始內容（配分 5%）")
    add_body("CH2 行程概念（Process Concept / Five-State Model 行程狀態轉換圖）。", bold_prefix="考點定位：")
    add_heading_3("架構模型圖示與狀態定義：")
    add_code([
        "               [Interrupt / Time Slice Exceeded]",
        "                      +-------------------+",
        "                      |                   |",
        "                      v                   |",
        "  +-------+  admit  +-------+  dispatch +---------+  exit  +------------+",
        "  |  New  | ------> | Ready | --------> | Running | -----> | Terminated |",
        "  +-------+         +-------+           +---------+        +------------+",
        "                      ^                   |",
        "                      |   I/O or Event    |",
        "                      |    Completion     |",
        "                      +-------------------+",
        "                                |",
        "                                v",
        "                        +---------------+",
        "                        | Waiting/Block |",
        "                        +---------------+"
    ])
    add_bullet("New（新建）：行程正在被建立中，尚未完全載入記憶體分配 PCB。", bold_prefix="1. 五大狀態精確定義：")
    add_bullet("Ready（就緒）：行程已具備執行所需之全部資源，僅等待 CPU 排程器分派處理器時間。", level=1)
    add_bullet("Running（執行）：行程指令正在 CPU 上執行。", level=1)
    add_bullet("Waiting / Blocked（等待 / 阻塞）：行程因等待特定事件（如 I/O 完成或收到信號）而暫停，此時即使 CPU 閒置亦無法執行。", level=1)
    add_bullet("Terminated（結束）：行程已執行完畢或被強制中止，作業系統正在回收其系統資源。", level=1)
    add_bullet("New -> Ready：經長程排程器核准（Admitted）載入就緒佇列。", bold_prefix="2. 六大狀態轉移觸發條件：")
    add_bullet("Ready -> Running：經短程排程器分派（Scheduler Dispatch）獲得 CPU。", level=1)
    add_bullet("Running -> Ready：時鐘中斷、時間配額用盡（Time Slice Expired）或遭更高優先權行程搶佔。", level=1)
    add_bullet("Running -> Waiting：發起 I/O 請求或等待事件同步（I/O or Event Wait）。", level=1)
    add_bullet("Waiting -> Ready：所等待之 I/O 作業完成或事件發生（I/O or Event Completion）。", level=1)
    add_bullet("Running -> Terminated：程式正常執行 exit() 系統呼叫或遭遇不可回復之例外錯誤。", level=1)

    # --- 考題 5: 109 年第六題 ---
    add_heading_2("考題五、【109 學年度】第六題：死結發生四大必要條件（Four Conditions for Deadlock）")
    add_callout([
        "第六題: (4 分) What is four conditions of deadlock occurring?"
    ], title="試題原始內容（配分 4%）")
    add_body("CH5 死結特性（Deadlock Characterization / Coffman Conditions）。", bold_prefix="考點定位：")
    add_heading_3("四大必要條件詳解：")
    add_bullet("至少有一項資源必須處於非共享模式，亦即一次只能由一個行程持有與使用；若有其他行程申請該資源，申請者必須等待至資源被釋放為止。", bold_prefix="1. 互斥條件（Mutual Exclusion）：")
    add_bullet("一個行程必須同時持有至少一項資源，並且正在等待取得其他被其他行程所持有的額外資源。", bold_prefix="2. 持有並等待（Hold and Wait）：")
    add_bullet("資源無法被強制剝奪；資源只能在持有它的行程完成任務後，由該行程自願且主動釋放。", bold_prefix="3. 不可搶佔（No Preemption）：")
    add_bullet("系統中存在一組封閉的等待行程序列 {P0, P1, ..., Pn}，其中 P0 正等待 P1 持有的資源，P1 正等待 P2 持有的資源，...，Pn 正等待 P0 持有的資源，形成循環閉環相依等待鏈。", bold_prefix="4. 循環等待（Circular Wait）：")
    add_body("特別備註：上述四大條件必須「同時成立」死結方會發生；死結預防（Deadlock Prevention）即奠基於藉由設計協定破壞上述四大條件中的至少一項。", bold_prefix="理論閉環：")

    # --- 考題 6: 109 年第八題 ---
    add_heading_2("考題六、【109 學年度】第八題：即時 CPU 排程（RMS vs. EDF）充足性定理與推導")
    add_callout([
        "第八題: (6 分) Without any doubt, real-time scheduler must support preemptive and priority-based scheduling. For hard real-time scheduling must provide ability to meet deadlines. Given two processes P1 and P2, the periodic generated time of P1 and P2 are p1=50 and p2=80, respectively. And the processing time of P1 and P2 are t1=25 and t2=35, respectively.",
        "(A) Can these two processes be scheduled using rate-monotonic scheduling? Why?",
        "(B) Can these two processes be scheduled using earliest-deadline-first scheduling (EDF)? Why?"
    ], title="試題原始內容（配分 6%）")
    add_body("CH3 CPU 排程（即時排程演算法 Real-Time Scheduling：RMS 速率單調 vs. EDF 最早截止時間優先）。", bold_prefix="考點定位：")
    add_heading_3("參數梳理與系統利用率結算：")
    add_bullet("任務 P1：週期 p1 = 50，執行時間 t1 = 25，利用率 U1 = 25/50 = 0.5。")
    add_bullet("任務 P2：週期 p2 = 80，執行時間 t2 = 35，利用率 U2 = 35/80 = 0.4375。")
    add_bullet("系統總 CPU 利用率：U = U1 + U2 = 0.5 + 0.4375 = 0.9375（即 93.75%）。")

    add_heading_3("小題 (A) Rate-Monotonic Scheduling (RMS) 分析：")
    add_bullet("優先權指派：RMS 為靜態優先權，週期越短優先權越高。因 p1 = 50 < p2 = 80，故優先權 P1 > P2。")
    add_bullet("Liu & Layland 充足條件檢驗：U <= N * (2^(1/N) - 1)。當 N = 2 時，U_bound = 2 * (sqrt(2) - 1) ≈ 82.84%。因總利用率 U = 93.75% > 82.84%，超過充分條件上限，必須透過時間軸排程（甘特圖）精確驗證。")
    add_bullet("甘特圖時間軸模擬：")
    add_bullet("t = 0 ~ 25：P1 執行並完成（Deadline = 50）。", level=1)
    add_bullet("t = 25 ~ 50：P2 開始執行，執行了 50 - 25 = 25 單位（尚餘 35 - 25 = 10 單位）。", level=1)
    add_bullet("t = 50：P1 第二個週期到達，因優先權較高，強制搶佔 CPU。", level=1)
    add_bullet("t = 50 ~ 75：P1 執行第二個週期並於 t = 75 完成。", level=1)
    add_bullet("t = 75：CPU 交還給 P2，P2 尚需 10 單位運算，預計於 t = 75 + 10 = 85 方能完成。", level=1)
    add_bullet("違規判定：P2 第一個週期的絕對截止時間為 t = 80！因 85 > 80，P2 錯過了截止時間（Missed Deadline）！", level=1)
    add_body("結論 (A)：不可以（No）。因為總利用率 93.75% 超過 RMS 充分保證上限 82.84%，且甘特圖實證推導證明在 t = 80 時 P2 違反截止時間。", bold_prefix="判定結論：")

    add_heading_3("小題 (B) Earliest-Deadline-First Scheduling (EDF) 分析：")
    add_bullet("排程核心原理：EDF 採動態優先權，誰的絕對截止時間（Absolute Deadline）越近，誰的優先權越高。")
    add_bullet("充要條件定理：對於任意週期性任務集合，只要滿足總利用率 U = sum(t_i / p_i) <= 1.0 (100%)，EDF 保證所有任務皆能滿足截止時間（Schedulable）。")
    add_bullet("本系統總利用率 U = 0.9375 <= 1.0，完全符合充要定理保證。")
    add_bullet("EDF 動態時間軸驗證：")
    add_bullet("t = 0 ~ 25：P1（Deadline 50）比 P2（Deadline 80）早，P1 執行至 25 完成。", level=1)
    add_bullet("t = 25 ~ 50：P2 執行 25 單位（剩 10 單位）。", level=1)
    add_bullet("t = 50：P1 第二週期到達，其 Deadline 為 50 + 50 = 100。此時 P2 的 Deadline 為 80。因 80 < 100，P2 的截止時間更早，故 P2 不會被 P1 搶佔！", level=1)
    add_bullet("t = 50 ~ 60：P2 繼續執行完畢（t = 60 <= 80，成功達標）。", level=1)
    add_bullet("t = 60 ~ 85：P1 執行第二週期完畢（t = 85 <= 100，成功達標）。後續週期循環皆能準時完成。", level=1)
    add_body("結論 (B)：可以（Yes）。因為 EDF 排程之充分且必要條件為總利用率 U <= 1.0，本系統之 93.75% <= 100%，因此保證可成功排程且絕不違反截止時間。", bold_prefix="判定結論：")

    # --- 考題 7: 110 年第六題 ---
    add_heading_2("考題七、【110 學年度】第六題：Peterson 演算法程式填空與互斥邏輯")
    add_callout([
        "六、(10%) To use Peterson's solution to solve the critical-section problem between two processes Pi and Pj, please fill the blank (a), (b), (c), (d) and (e) in structure of Pi in Program 1.",
        "do {",
        "    flag[i] = _(a)_ ;",
        "    turn = _(b)_ ;",
        "    while ( _(c)_ && _(d)_ );",
        "        critical section",
        "    flag[i] = _(e)_ ;",
        "        remainder section",
        "} while (true);"
    ], title="試題原始內容（配分 10%）")
    add_body("CH2 / CH6 行程同步（Process Synchronization / Peterson's Algorithm 經典雙行程軟體解法）。", bold_prefix="考點定位：")
    add_heading_3("完整詳解與變數邏輯分析：")
    add_bullet("flag[i]：布林陣列，表示行程 Pi 是否具備進入臨界區間之意圖。Pi 進入前先宣告 flag[i] = true。離開時重設為 false。")
    add_bullet("turn：整數變數，表示當前禮讓給哪一個行程優先進入。Pi 會主動禮讓對方，將 turn = j。")
    add_bullet("忙碌等待迴圈：while (flag[j] && turn == j);。唯有當對方「也想進去（flag[j] == true）」且「輪到對方（turn == j）」時，Pi 才會在迴圈中自旋等待；一旦對方離開將 flag[j] 設為 false，或對方把 turn 讓回給 i，Pi 即可進入臨界區間。")
    add_bullet("(a) true", bold_prefix="標準填空答案：")
    add_bullet("(b) j", bold_prefix="")
    add_bullet("(c) flag[j]（可與 (d) 對調）", bold_prefix="")
    add_bullet("(d) turn == j（可與 (c) 對調）", bold_prefix="")
    add_bullet("(e) false", bold_prefix="")

    # --- 考題 8: 111 年第 3 題 ---
    add_heading_2("考題八、【111 學年度】第 3 題：多重實例資源死結偵測演算法逐步推導")
    add_callout([
        "3. (5%) 假設一系統包含有四種硬體資源 A、B、C、D，現有五個行程(process) P1、P2、P3、P4、P5 在此系統中，假設系統目前資源分配狀態如下：",
        "   Allocation: P1(1,0,0,1), P2(1,1,1,1), P3(1,1,0,1), P4(1,1,0,0), P5(1,1,1,1)",
        "   Request:    P1(1,0,0,1), P2(0,1,1,1), P3(1,0,1,0), P4(1,1,0,1), P5(0,1,1,0)",
        "   Available:  (1, 1, 0, 1)",
        "   (a) 請問每種硬體資源分別有多少實體(instance)？",
        "   (b) 請解釋此系統目前是否存在死結(deadlock)？如有死結，是哪幾個行程所造成？"
    ], title="試題原始內容（配分 5%）")
    add_body("CH5 死結偵測（Deadlock Detection Algorithm 多重實例矩陣推導）。", bold_prefix="考點定位：")
    add_heading_3("逐步計算推導歷程：")
    add_body("小題 (a) 求資源總實體數（Total Instances）：")
    add_bullet("公式：Total = Available + sum(Allocation_i)")
    add_bullet("各資源已分配總和：")
    add_bullet("A: 1 + 1 + 1 + 1 + 1 = 5", level=1)
    add_bullet("B: 0 + 1 + 1 + 1 + 1 = 4", level=1)
    add_bullet("C: 0 + 1 + 0 + 0 + 1 = 2", level=1)
    add_bullet("D: 1 + 1 + 1 + 0 + 1 = 4", level=1)
    add_bullet("加總目前 Available = (1, 1, 0, 1)：")
    add_bullet("Total A = 5 + 1 = 6 個實體", level=1)
    add_bullet("Total B = 4 + 1 = 5 個實體", level=1)
    add_bullet("Total C = 2 + 0 = 2 個實體", level=1)
    add_bullet("Total D = 4 + 1 = 5 個實體", level=1)
    add_body("答案 (a)：A 資源 6 個、B 資源 5 個、C 資源 2 個、D 資源 5 個。", bold_prefix="小題 (a) 結論：")

    add_body("小題 (b) 死結偵測演算法推導：")
    add_bullet("初始化：令 Work = Available = (1, 1, 0, 1)，Finish 向量全為 False。")
    add_bullet("步驟 1：比對各行程之 Request_i <= Work (1, 1, 0, 1)：")
    add_bullet("P1: Request_1 (1, 0, 0, 1) <= Work (1, 1, 0, 1) 滿足！P1 執行完畢並釋放資源：Work = (1, 1, 0, 1) + Allocation_1 (1, 0, 0, 1) = (2, 1, 0, 2)，Finish[1] = True。", level=1)
    add_bullet("步驟 2：目前 Work = (2, 1, 0, 2)，比對剩餘行程之 Request：")
    add_bullet("P2: Request_2 = (0, 1, 1, 1)，需 C=1，但 Work_C = 0，無法滿足。", level=1)
    add_bullet("P3: Request_3 = (1, 0, 1, 0)，需 C=1，但 Work_C = 0，無法滿足。", level=1)
    add_bullet("P4: Request_4 = (1, 1, 0, 1) <= Work (2, 1, 0, 2) 滿足！P4 執行完畢並釋放資源：Work = (2, 1, 0, 2) + Allocation_4 (1, 1, 0, 0) = (3, 2, 0, 2)，Finish[4] = True。", level=1)
    add_bullet("步驟 3：目前 Work = (3, 2, 0, 2)，比對剩餘行程 P2, P3, P5：")
    add_bullet("P2: 需 C=1，無法滿足。", level=1)
    add_bullet("P3: 需 C=1，無法滿足。", level=1)
    add_bullet("P5: Request_5 = (0, 1, 1, 0)，需 C=1，無法滿足。", level=1)
    add_bullet("步驟 4：此時系統已無其他行程可釋放資源，但 P2、P3、P5 皆因缺乏資源 C（目前可用數為 0）而永久處於阻塞等待狀態。")
    add_body("答案 (b)：系統目前「存在死結（Deadlock exists）」，陷入死結的行程為 P2、P3、P5。", bold_prefix="小題 (b) 結論：")

    # --- 考題 9: 113 年第 4 題 ---
    add_heading_2("考題九、【113 學年度】第 4 題：fork() 與 thread_create() 行程與執行緒計數推導")
    add_callout([
        "4. (5%) 考慮一程式碼區段：",
        "   pid_t pid;",
        "   pid = fork();",
        "   if (pid == 0) { /* child process */",
        "       fork();",
        "       thread_create( ... );",
        "   }",
        "   fork();",
        "   甲、請問會產生多少個獨立的行程(process)？(2%)",
        "   乙、請問會產生多少個獨立的執行緒(thread)？(3%)"
    ], title="試題原始內容（配分 5%）")
    add_body("CH2 行程建立（fork() 行程樹）與 CH4 多執行緒程式設計（執行緒建立與 POSIX fork 語意）。", bold_prefix="考點定位：")
    add_heading_3("逐步追蹤推導：")
    add_bullet("初始狀態：系統存在 1 個初始父行程 P0，內部包含 1 個主執行緒。")
    add_bullet("第 1 行 pid = fork();：P0 呼叫 fork()，分裂產生子行程 P1。此時有 2 個行程：P0（pid > 0，父行程）與 P1（pid == 0，子行程）。")
    add_bullet("進入 if 區塊：只有 pid == 0 的子行程 P1 進入 if 內部；P0 跳過。")
    add_bullet("在 if 內部：P1 呼叫 fork() 產生子行程 P2。此時 if 內有 P1 與 P2 兩個行程。接著 P1 與 P2 各自呼叫一次 thread_create(...)，各建立 1 個新執行緒。")
    add_bullet("離開 if 區塊時系統共有 3 個行程：")
    add_bullet("P0：擁有 1 個主執行緒。", level=1)
    add_bullet("P1：擁有 1 個主執行緒 + 1 個新建執行緒 = 2 個執行緒。", level=1)
    add_bullet("P2：擁有 1 個主執行緒 + 1 個新建執行緒 = 2 個執行緒。", level=1)
    add_bullet("最後一行 fork();：目前存在的 3 個行程（P0, P1, P2）各自呼叫一次 fork()！行程總數直接翻倍：3 * 2 = 6 個行程。")
    add_bullet("依據 POSIX 規範，在多執行緒環境中呼叫 fork() 時，子行程僅複製呼叫 fork() 的單一執行緒，其餘執行緒不會被複製至子行程中：")
    add_bullet("P0（1 執行緒）fork 產生 P0'（1 執行緒）。", level=1)
    add_bullet("P1（2 執行緒，主執行緒呼叫 fork）fork 產生 P1'（1 執行緒）。", level=1)
    add_bullet("P2（2 執行緒，主執行緒呼叫 fork）fork 產生 P2'（1 執行緒）。", level=1)
    add_bullet("執行緒總數結算：1(P0) + 1(P0') + 2(P1) + 1(P1') + 2(P2) + 1(P2') = 8 個執行緒。")
    add_body("甲、產生 6 個獨立行程（若問新建立則為 5 個）。\n乙、產生 8 個獨立執行緒（若按抽象模型假設子行程複製全部執行緒則為 10 個，標準 POSIX 為 8 個）。", bold_prefix="標準答案：")

    # --- 考題 10: 114 年第 12 題 ---
    add_heading_2("考題十、【114 學年度】第 12 題：多層佇列 vs. 多層反饋佇列核心差異")
    add_callout([
        "12. What is the main difference between the \"Multilevel Queue\" and \"Multilevel Feedback Queue\" scheduling algorithms?",
        "(A) process migration among queues",
        "(B) number of queues",
        "(C) time quantum",
        "(D) level definition",
        "(E) size of queue"
    ], title="試題原始內容（單選題，配分 2%）")
    add_body("CH3 CPU 排程演算法（Multilevel Queue vs. Multilevel Feedback Queue）。", bold_prefix="考點定位：")
    add_body("(A) process migration among queues", bold_prefix="正確答案：")
    add_bullet("Multilevel Queue（多層佇列排程）：行程在進入系統時被永久分類指派至特定佇列（如前景互動式或背景批次式），行程在整個生命週期中「無法在不同佇列之間移動（No process migration）」。", bold_prefix="選項解析：")
    add_bullet("Multilevel Feedback Queue（多層反饋佇列排程，MLFQ）：允許行程依據其 CPU Burst 行為在不同佇列之間動態轉移（Process migration among queues）。消耗過多 CPU 時間的行程會被降級至低優先權佇列，長時間未執行的行程可藉由老化（Aging）機制升級至高優先權佇列以防飢餓。")

    # --- 考題 11: 114 年第 14 題 ---
    add_heading_2("考題十一、【114 學年度】第 14 題：巢狀 fork() 與 printf 輸出次數執行路徑分析")
    add_callout([
        "14. How many times does the following program print hello?",
        "#include <stdio.h>",
        "#include <unistd.h>",
        "main () {",
        "    if ( fork () == 0)",
        "        printf(\"hello\\n\");",
        "    else if ( fork () == 0)",
        "        printf(\"hello\\n\");",
        "    printf(\"hello\\n\");",
        "}",
        "(A) 2  (B) 3  (C) 4  (D) 5  (E) 6"
    ], title="試題原始內容（單選題，配分 2%）")
    add_body("CH2 行程概念與 fork() 條件分支執行流程。", bold_prefix="考點定位：")
    add_body("(D) 5", bold_prefix="正確答案：")
    add_bullet("設初始父行程為 P0。")
    add_bullet("執行 if ( fork() == 0 )：P0 分裂出子行程 C1。")
    add_bullet("對於子行程 C1：fork() 回傳 0，條件成立，進入 if 內部印出第 1 次 \"hello\\n\"。跳過 else if，最後執行末尾的 printf(\"hello\\n\");，印出第 2 次 \"hello\\n\"。C1 累計印出 2 次。", level=1)
    add_bullet("對於父行程 P0：fork() 回傳非 0，條件為假，進入 else if ( fork() == 0 ) 分支。", level=1)
    add_bullet("在 else if 中 P0 再次呼叫 fork()，分裂出子行程 C2：")
    add_bullet("對於子行程 C2：fork() 回傳 0，條件成立，進入 else if 內部印出第 3 次 \"hello\\n\"。隨後執行末尾的 printf(\"hello\\n\");，印出第 4 次 \"hello\\n\"。C2 累計印出 2 次。", level=1)
    add_bullet("對於父行程 P0：fork() 回傳非 0，條件為假，不進入 else if 內部。直接執行末尾的 printf(\"hello\\n\");，印出第 5 次 \"hello\\n\"。P0 累計印出 1 次。", level=1)
    add_bullet("全系統印出 hello 總次數：2 (C1) + 2 (C2) + 1 (P0) = 5 次。")

    # --- 考題 12: 114 年第 29 題 ---
    add_heading_2("考題十二、【114 學年度】第 29 題：行程控制表（PCB）收錄資訊項目")
    add_callout([
        "29. Which following information is stored in process control block (PCB) of a process? (複選)",
        "(A) scheduling information",
        "(B) accounting information",
        "(C) inter communication information",
        "(D) I/O status information",
        "(E) process state information"
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH2 行程概念（Process Control Block, PCB 資料結構）。", bold_prefix="考點定位：")
    add_body("(A), (B), (C), (D), (E)（全選）", bold_prefix="正確答案：")
    add_bullet("(A) scheduling information：正確。包含行程優先權、排程佇列指標與排程參數。", bold_prefix="選項解析：")
    add_bullet("(B) accounting information：正確。包含 CPU 使用時間累積量、時限上限與帳號資料。")
    add_bullet("(C) inter communication information：正確。包含行程間通訊之旗標、訊號處理表與訊息佇列指標。")
    add_bullet("(D) I/O status information：正確。包含配置給該行程的 I/O 裝置清單與開啟檔案描述元表格。")
    add_bullet("(E) process state information：正確。包含行程當前狀態（New, Ready, Running, Waiting, Terminated）。")

    # --- 考題 13: 114 年第 31 題 ---
    add_heading_2("考題十三、【114 學年度】第 31 題：臨界區間問題三大必要準則")
    add_callout([
        "31. A solution to the critical-section problem must satisfy which requirements? (複選)",
        "(A) mutual exclusion",
        "(B) circular wait",
        "(C) progress",
        "(D) no preemption",
        "(E) bounded waiting"
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH2 / CH6 行程同步（Critical-Section Problem 三大要件）。", bold_prefix="考點定位：")
    add_body("(A), (C), (E)", bold_prefix="正確答案：")
    add_bullet("(A) mutual exclusion（互斥）：必須保證任一時刻最多只有一個行程在臨界區間執行。", bold_prefix="選項解析：")
    add_bullet("(C) progress（進展）：當臨界區間空閒且有行程欲進入時，只有不在剩餘區間執行的行程可參與決策，且不可無限期推遲。", bold_prefix="")
    add_bullet("(E) bounded waiting（有限等待）：行程提出請求至獲准進入之間，其他行程獲准進入的次數必須有界限，防止飢餓。", bold_prefix="")
    add_bullet("選項 (B) circular wait 與 (D) no preemption 為死結（Deadlock）之必要條件，並非臨界區間的解法要件。", bold_prefix="干擾項排除：")

    # --- 考題 14: 114 年第 33 題 ---
    add_heading_2("考題十四、【114 學年度】第 33 題：死結與活結特性及防範手段分析")
    add_callout([
        "33. What are the correct descriptions of deadlock and livelock? (複選)",
        "(A) Processes in deadlock can not proceed until the other processes releases the resource",
        "(B) A livelock exists if the processes are changing their state infrequently and proceeding ineffectively.",
        "(C) Processes in livelock constantly change their state in response to each other, but none of them can make any progress",
        "(D) One way to prevent a deadlock is to use a resource ordering protocol for assigning a unique priority to each resource",
        "(E) One way to detect a livelock is to monitor and reorder the states of the processes over time"
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH5 死結與活結特性（Deadlock & Livelock 原理與預防機制）。", bold_prefix="考點定位：")
    add_body("(A), (C), (D)", bold_prefix="正確答案：")
    add_bullet("(A) 正確：死結行程處於永久等待狀態，若無其他持有者釋放資源，則永久無法推進。", bold_prefix="選項解析：")
    add_bullet("(B) 錯誤：活結行程並非 \"infrequently\" 改變狀態，而是極度頻繁（constantly / actively）相互回應改變自身狀態。")
    add_bullet("(C) 正確：活結標準定義為行程不斷對彼此狀態做出反應並頻繁切換狀態，但無任何行程能取得實質進展。")
    add_bullet("(D) 正確：破壞循環等待的標準預防機制即為全域資源遞增排序（Resource Ordering Protocol），強制所有行程依序申請。")
    add_bullet("(E) 錯誤：活結行程並非處於阻塞佇列，CPU 使用率往往極高，單純監控狀態重排難以有效偵測活結。")

    # --- 考題 15: 115 年第 21 題 ---
    add_heading_2("考題十五、【115 學年度】第 21 題：死結避免、活結代價與系統吞吐量全維度研判")
    add_callout([
        "21. Which following statements are correct about deadlock and livelock? (複選)",
        "(A) Breaking the hold-and-wait condition could prevent deadlock, but may lead to livelock.",
        "(B) The Banker's Algorithm could prevent livelock directly.",
        "(C) Enforcing a strict resource ordering policy could reduce the risk of both deadlock and livelock occurrence.",
        "(D) Aggressive deadlock avoidance strategies may reduce overall system throughput.",
        "(E) Livelock is generally easier to detect than deadlock through system monitoring tools in OS."
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH5 死結預防、避免與活結代價分析（Deadlock vs. Livelock）。", bold_prefix="考點定位：")
    add_body("(A), (C), (D)", bold_prefix="正確答案：")
    add_bullet("(A) 正確：打破 Hold-and-Wait 的做法常要求「若無法取得全部資源，則釋放所有已持有資源並重新嘗試」；若多個行程反覆同時釋放並重新爭搶，將陷入活結震盪。", bold_prefix="選項解析：")
    add_bullet("(B) 錯誤：銀行家演算法是死結避免演算法，專門用於迴避不安全狀態以防死結，無法直接解決演算法重試邏輯造成的活結。")
    add_bullet("(C) 正確：嚴格的資源排序協定強制所有行程循相同方向申請資源，徹底消除循環等待，同時大幅降低行程互相禮讓震盪引發活結之機率。")
    add_bullet("(D) 正確：過度嚴格保守的死結避免策略會推遲許多實際不會引發死結的資源請求，導致資源閒置並降低系統整體吞吐量。")
    add_bullet("(E) 錯誤：活結行程處於 Running 狀態而非 Waiting 狀態，傳統 OS 監控工具難以將其與正常高效運算區分，因此活結遠比死結更難偵測。")

    # --- 考題 16: 115 年第 22 題 ---
    add_heading_2("考題十六、【115 學年度】第 22 題：硬體同步機制與臨界區間互斥保證")
    add_callout([
        "22. Regarding hardware-based solutions to the Critical-Section Problem, which of the following statements are correct? (複選)",
        "(A) Hardware atomic instructions eliminate busy waiting.",
        "(B) Disabling interrupts is suitable for multiprocessor systems.",
        "(C) Disabling interrupts guarantees mutual exclusion on a uniprocessor system.",
        "(D) Test-and-Set instructions can be used to implement spinlocks.",
        "(E) Compare-and-Swap (CAS) can be used to build mutual exclusion primitives."
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH2 / CH6 行程同步之硬體支援（Hardware-Based Synchronization Primitives）。", bold_prefix="考點定位：")
    add_body("(C), (D), (E)", bold_prefix="正確答案：")
    add_bullet("(A) 錯誤：硬體原子指令（如 Test-and-Set）僅提供不可分割的讀寫原語，行程在等待鎖時依然需透過自旋迴圈 while(test_and_set(&lock)) 進行忙碌等待，並未消除 busy waiting。", bold_prefix="選項解析：")
    add_bullet("(B) 錯誤：在多處理器系統（Multiprocessor）中關閉中斷極為低效且不可行，因為必須向所有處理器發送跨核心中斷關閉訊號，開銷過於龐大。")
    add_bullet("(C) 正確：在單處理器系統（Uniprocessor）中，只要關閉中斷，CPU 就不會發生時鐘中斷與上下文切換，當前行程即可不被打斷地執行完臨界區間，完全保證互斥。")
    add_bullet("(D) 正確：test_and_set 為實作自旋鎖（Spinlocks）最經典的硬體原子指令。")
    add_bullet("(E) 正確：Compare-and-Swap (CAS) 為現代處理器（如 x86 CMPXCHG）構建互斥鎖、號誌與 Lock-free 演算法的核心原語。")

    # --- 考題 17: 115 年第 25 題 ---
    add_heading_2("考題十七、【115 學年度】第 25 題：即時 CPU 排程（RMS 與 EDF 特性研判）")
    add_callout([
        "25. Which of the following statements about Real-Time CPU Scheduling are correct? (複選)",
        "(A) Hard real-time systems cannot tolerate missed deadlines.",
        "(B) Rate Monotonic Scheduling assigns priority based on period length.",
        "(C) Soft real-time systems guarantee deadlines.",
        "(D) Earliest Deadline First Scheduling dynamically adjusts priorities.",
        "(E) Earliest Deadline First Scheduling could theoretically achieve 100% CPU utilization."
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH3 CPU 排程（即時排程演算法 RMS 與 EDF 原理）。", bold_prefix="考點定位：")
    add_body("(A), (B), (D), (E)", bold_prefix="正確答案：")
    add_bullet("(A) 正確：硬即時系統（Hard real-time）要求關鍵工作必須在嚴格截止時間內完成，錯過截止時間即判定為全系統致命故障，不可容忍任何逾時。", bold_prefix="選項解析：")
    add_bullet("(B) 正確：RMS 為靜態優先權演算法，依據任務週期長度指派優先權（週期越短優先權越高）。")
    add_bullet("(C) 錯誤：軟即時系統（Soft real-time）僅提供盡力保證（Best effort），優先執行即時工作，但無法保證絕對不違背截止時間。")
    add_bullet("(D) 正確：EDF 依據任務當前絕對截止時間動態調整優先權（截止時間越近優先權越高）。")
    add_bullet("(E) 正確：在理論上只要週期性任務總 CPU 利用率 <= 100%，EDF 保證所有任務皆能滿足截止時間。")

    # --- 考題 18: 115 年第 27 題 ---
    add_heading_2("考題十八、【115 學年度】第 27 題：行程與執行緒同步、記憶體共享與 IPC 開銷")
    add_callout([
        "27. To tackle the synchronization and communication of processes and threads, which of the following statements are correct? (複選)",
        "(A) Different processes could directly share each other's local (stack) variables by default.",
        "(B) Threads within the same process could communicate through shared memory.",
        "(C) Mutexes and semaphores could be used for thread synchronization.",
        "(D) Inter-process communication (IPC) usually has higher overhead than communication between threads within the same process.",
        "(E) Threads within the same process will not cause race conditions."
    ], title="試題原始內容（複選題，配分 2%）")
    add_body("CH2 行程概念與 IPC、CH4 多執行緒記憶體架構（Threads vs. Processes）。", bold_prefix="考點定位：")
    add_body("(B), (C), (D)", bold_prefix="正確答案：")
    add_bullet("(A) 錯誤：預設情況下，不同行程各自擁有獨立受保護的虛擬定址空間，無法直接存取對方的區域堆疊（Stack）變數。", bold_prefix="選項解析：")
    add_bullet("(B) 正確：同一行程內部所屬的多個執行緒共享相同定址空間（包含 Data 段與 Heap 段），可直接透過共享記憶體傳遞資料。", bold_prefix="")
    add_bullet("(C) 正確：互斥鎖（Mutexes）與信號量（Semaphores）為執行緒同步的基礎原語。", bold_prefix="")
    add_bullet("(D) 正確：IPC 涉及作業系統核心模式切換與跨位址空間資料拷貝，系統開銷遠高於同一行程內執行緒直接讀寫記憶體。", bold_prefix="")
    add_bullet("(E) 錯誤：正因同一行程內的多執行緒共享記憶體，若缺乏適當同步保護，同時讀寫將直接引發競態條件（Race Conditions）。", bold_prefix="")

    # 儲存檔案
    output_filename = "中央大學研究所_作業系統歷年考題篩選與期中範圍詳解.docx"
    doc.save(output_filename)
    print(f"File successfully created: {os.path.abspath(output_filename)}")

if __name__ == "__main__":
    build_document()
